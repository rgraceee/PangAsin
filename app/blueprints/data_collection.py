from flask import Blueprint, render_template, redirect, url_for, flash, request
from flask_login import login_required, current_user
from app.forms.production import ProductionRecordForm
from app.forms.upload import UploadForm
from app.models.production_record import ProductionRecord
from app.models.submission_batch import SubmissionBatch
from app.models.validation_history import ValidationHistory
from app.extensions import db
from datetime import datetime
import openpyxl

encoder_bp = Blueprint("encoder", __name__)


def encoder_required():
    if not current_user.is_authenticated or current_user.role != "encoder":
        flash("Encoder access required.", "danger")
        return False
    return True


@encoder_bp.before_request
def require_encoder():
    if not encoder_required():
        return redirect(url_for("monitoring.landing"))


@encoder_bp.route("/encoder/home")
@login_required
def home():
    records = ProductionRecord.query.filter_by(submitted_by=current_user.id).order_by(ProductionRecord.created_at.desc()).all()
    pending = sum(1 for r in records if r.status == "pending")
    approved = sum(1 for r in records if r.status == "approved")
    rejected = sum(1 for r in records if r.status == "rejected")
    return render_template("encoder/home.html", records=records, pending=pending, approved=approved, rejected=rejected)


@encoder_bp.route("/encoder/submit", methods=["GET", "POST"])
@login_required
def submit():
    form = ProductionRecordForm()
    if form.validate_on_submit():
        is_draft = form.save_draft.data and not form.submit.data
        record = ProductionRecord(
            municipality_id=current_user.municipality_id,
            submission_batch_id=None,
            barangay=form.barangay.data,
            period_type=form.period_type.data,
            record_date=form.record_date.data,
            production_volume=form.production_volume.data,
            production_method=form.production_method.data,
            production_area=form.production_area.data,
            num_salt_beds=form.num_salt_beds.data,
            producer_age=form.producer_age.data,
            producer_gender=form.producer_gender.data,
            notes=form.notes.data,
            status="draft" if is_draft else "pending",
            submitted_by=current_user.id,
            submitted_at=datetime.utcnow() if not is_draft else None,
        )
        db.session.add(record)
        db.session.commit()
        flash("Record saved." if is_draft else "Record submitted for validation.", "success")
        return redirect(url_for("encoder.home"))
    return render_template("encoder/submit_form.html", form=form, title="Submit Record")


@encoder_bp.route("/encoder/records/<int:id>/edit", methods=["GET", "POST"])
@login_required
def edit_record(id):
    record = ProductionRecord.query.get_or_404(id)
    if record.submitted_by != current_user.id or record.status not in ("draft", "rejected"):
        flash("You cannot edit this record.", "danger")
        return redirect(url_for("encoder.submissions"))
    form = ProductionRecordForm(obj=record)
    if form.validate_on_submit():
        previous_status = record.status
        record.period_type = form.period_type.data
        record.record_date = form.record_date.data
        record.barangay = form.barangay.data
        record.production_volume = form.production_volume.data
        record.production_method = form.production_method.data
        record.production_area = form.production_area.data
        record.num_salt_beds = form.num_salt_beds.data
        record.producer_age = form.producer_age.data
        record.producer_gender = form.producer_gender.data
        record.notes = form.notes.data
        record.status = "pending"
        record.submitted_at = datetime.utcnow()
        db.session.commit()
        if previous_status == "rejected":
            vh = ValidationHistory(
                production_record_id=record.id,
                previous_status=previous_status,
                new_status="pending",
                action="resubmitted",
                reviewer_id=None,
                comment=None,
            )
            db.session.add(vh)
            db.session.commit()
        flash("Record updated and resubmitted for validation." if previous_status == "rejected" else "Record updated.", "success")
        return redirect(url_for("encoder.submissions"))
    return render_template("encoder/submit_form.html", form=form, title="Edit Record", record=record)


@encoder_bp.route("/encoder/submissions")
@login_required
def submissions():
    status_filter = request.args.get("status")
    query = ProductionRecord.query.filter_by(submitted_by=current_user.id)
    if status_filter:
        query = query.filter_by(status=status_filter)
    records = query.order_by(ProductionRecord.created_at.desc()).all()
    return render_template("encoder/submissions.html", records=records, status_filter=status_filter)


@encoder_bp.route("/encoder/upload", methods=["GET", "POST"])
@login_required
def upload():
    form = UploadForm()
    if form.validate_on_submit():
        file = form.file.data
        filename = file.filename
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        valid_rows = []
        invalid_rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            errors = []
            if not data.get("barangay"):
                errors.append("Missing barangay")
            if not data.get("period_type") or data.get("period_type") not in ("daily", "monthly", "annual"):
                errors.append("Invalid period_type")
            if not data.get("record_date"):
                errors.append("Missing record_date")
            try:
                vol = float(data.get("production_volume", 0))
                if vol < 0:
                    errors.append("Negative production_volume")
            except (TypeError, ValueError):
                errors.append("Invalid production_volume")
            if data.get("production_method") not in ("solar_evaporation", "cooked", "hybrid"):
                errors.append("Invalid production_method")
            if errors:
                invalid_rows.append({"row": row, "errors": errors})
            else:
                valid_rows.append(data)

        batch = SubmissionBatch(
            municipality_id=current_user.municipality_id,
            uploaded_by=current_user.id,
            original_filename=filename,
            total_rows=ws.max_row - 1,
            valid_rows=len(valid_rows),
            invalid_rows=len(invalid_rows),
            status="completed",
            processed_at=datetime.utcnow(),
        )
        db.session.add(batch)
        db.session.flush()

        for data in valid_rows:
            record = ProductionRecord(
                municipality_id=current_user.municipality_id,
                submission_batch_id=batch.id,
                barangay=data.get("barangay"),
                period_type=data.get("period_type"),
                record_date=data.get("record_date"),
                production_volume=data.get("production_volume"),
                production_method=data.get("production_method"),
                production_area=data.get("production_area"),
                num_salt_beds=data.get("num_salt_beds"),
                producer_age=data.get("producer_age"),
                producer_gender=data.get("producer_gender"),
                notes=data.get("notes"),
                status="pending",
                submitted_by=current_user.id,
                submitted_at=datetime.utcnow(),
            )
            db.session.add(record)

        db.session.commit()
        flash(f"Upload complete. Valid: {len(valid_rows)}, Invalid: {len(invalid_rows)}", "success")
        return redirect(url_for("encoder.home"))
    return render_template("encoder/upload.html", form=form)
