# WHAT: Report generation (build data per report type + write Excel/PDF/DOCX).
# WHY: Inilalayo ang file-format logic sa API para ma-reuse at madaling basahin.
import os
from datetime import datetime, date, timedelta

from sqlalchemy import func
from sqlalchemy.orm import joinedload

from app.extensions import db
from app.models.production_record import ProductionRecord
from app.models.municipality import Municipality
from app.models.barangay import Barangay
from app.models.user import User
from app.models.demand_benchmark import DemandBenchmark
from app.models.forecast import ForecastRun
from app.services.forecast_service import _monthly_aggregates
from app.constants import SECTOR_DEMAND


def _volume_mt(kg):
    return round((kg or 0) / 1000, 2)


def build_provincial_data(start, end):
    q = db.session.query(
        Municipality.id,
        Municipality.name,
        func.sum(ProductionRecord.production_volume).label("total_volume"),
        func.count(ProductionRecord.id).label("record_count"),
        func.sum(ProductionRecord.num_salt_beds).label("total_beds"),
        func.sum(ProductionRecord.num_salt_beds * ProductionRecord.area_per_salt_bed).label("total_area"),
        func.sum(ProductionRecord.registered_producers).label("total_registered"),
    ).join(ProductionRecord, ProductionRecord.municipality_id == Municipality.id)
    if start:
        q = q.filter(ProductionRecord.record_date >= start)
    if end:
        q = q.filter(ProductionRecord.record_date <= end)
    rows = q.group_by(Municipality.id, Municipality.name).all()

    labels, values, _ = _monthly_aggregates(None, start, end) if start and end else ([], [], [])

    return {
        "report_title": "Provincial Production Summary",
        "by_municipality": [
            {
                "municipality": r.name,
                "production_kg": round(float(r.total_volume or 0), 2),
                "production_mt": _volume_mt(r.total_volume),
                "record_count": r.record_count,
                "total_salt_beds": int(r.total_beds or 0),
                "total_area_sqm": round(float(r.total_area or 0), 2),
                "registered_producers": int(r.total_registered or 0),
            }
            for r in rows
        ],
        "monthly_trend": [
            {"month": labels[i], "production_mt": round(values[i] / 1000, 2)}
            for i in range(len(labels))
        ],
    }


def build_municipality_data(municipality_id, start, end):
    muni = db.session.get(Municipality, municipality_id)
    prefix = muni.name if muni else f"Municipality #{municipality_id}"

    q = db.session.query(
        Barangay.name.label("barangay"),
        func.sum(ProductionRecord.production_volume).label("total_volume"),
        func.count(ProductionRecord.id).label("record_count"),
        func.sum(ProductionRecord.num_salt_beds).label("total_beds"),
        func.sum(ProductionRecord.num_salt_beds * ProductionRecord.area_per_salt_bed).label("total_area"),
        func.sum(ProductionRecord.registered_producers).label("total_registered"),
    ).join(ProductionRecord, ProductionRecord.barangay_id == Barangay.id)
    if municipality_id:
        q = q.filter(ProductionRecord.municipality_id == municipality_id)
    if start:
        q = q.filter(ProductionRecord.record_date >= start)
    if end:
        q = q.filter(ProductionRecord.record_date <= end)
    rows = q.group_by(Barangay.name).all()

    total_volume = sum(float(r.total_volume or 0) for r in rows)

    return {
        "report_title": f"{prefix} Production Summary",
        "municipality_name": muni.name if muni else None,
        "by_barangay": [
            {
                "barangay": r.barangay,
                "production_kg": round(float(r.total_volume or 0), 2),
                "production_mt": _volume_mt(r.total_volume),
                "record_count": r.record_count,
                "total_salt_beds": int(r.total_beds or 0),
                "total_area_sqm": round(float(r.total_area or 0), 2),
                "registered_producers": int(r.total_registered or 0),
            }
            for r in rows
        ],
        "total_production_mt": _volume_mt(total_volume),
    }


def build_production_data(municipality_id, barangay_id, start, end):
    """Production-focused report across province, municipality, or barangay scope."""
    if barangay_id is not None:
        barangay = db.session.get(Barangay, barangay_id)
        prefix = barangay.name if barangay else f"Barangay #{barangay_id}"
        muni_name = barangay.municipality.name if barangay and barangay.municipality else None

        q = db.session.query(
            Barangay.name.label("barangay"),
            func.sum(ProductionRecord.production_volume).label("total_volume"),
            func.count(ProductionRecord.id).label("record_count"),
            func.sum(ProductionRecord.num_salt_beds).label("total_beds"),
            func.sum(ProductionRecord.num_salt_beds * ProductionRecord.area_per_salt_bed).label("total_area"),
        ).join(ProductionRecord, ProductionRecord.barangay_id == Barangay.id)
        q = q.filter(ProductionRecord.barangay_id == barangay_id)
        if start:
            q = q.filter(ProductionRecord.record_date >= start)
        if end:
            q = q.filter(ProductionRecord.record_date <= end)
        rows = q.group_by(Barangay.name).all()

        total_volume = sum(float(r.total_volume or 0) for r in rows)
        return {
            "report_title": f"{prefix} Production Summary",
            "municipality_name": muni_name,
            "by_barangay": [
                {
                    "barangay": r.barangay,
                    "production_kg": round(float(r.total_volume or 0), 2),
                    "production_mt": _volume_mt(r.total_volume),
                    "record_count": r.record_count,
                    "total_salt_beds": int(r.total_beds or 0),
                    "total_area_sqm": round(float(r.total_area or 0), 2),
                }
                for r in rows
            ],
            "total_production_mt": _volume_mt(total_volume),
        }

    if municipality_id is not None:
        muni = db.session.get(Municipality, municipality_id)
        prefix = muni.name if muni else f"Municipality #{municipality_id}"

        q = db.session.query(
            Barangay.name.label("barangay"),
            func.sum(ProductionRecord.production_volume).label("total_volume"),
            func.count(ProductionRecord.id).label("record_count"),
            func.sum(ProductionRecord.num_salt_beds).label("total_beds"),
            func.sum(ProductionRecord.num_salt_beds * ProductionRecord.area_per_salt_bed).label("total_area"),
        ).join(ProductionRecord, ProductionRecord.barangay_id == Barangay.id)
        q = q.filter(ProductionRecord.municipality_id == municipality_id)
        if start:
            q = q.filter(ProductionRecord.record_date >= start)
        if end:
            q = q.filter(ProductionRecord.record_date <= end)
        rows = q.group_by(Barangay.name).all()

        total_volume = sum(float(r.total_volume or 0) for r in rows)
        return {
            "report_title": f"{prefix} Production Summary",
            "municipality_name": muni.name if muni else None,
            "by_barangay": [
                {
                    "barangay": r.barangay,
                    "production_kg": round(float(r.total_volume or 0), 2),
                    "production_mt": _volume_mt(r.total_volume),
                    "record_count": r.record_count,
                    "total_salt_beds": int(r.total_beds or 0),
                    "total_area_sqm": round(float(r.total_area or 0), 2),
                }
                for r in rows
            ],
            "total_production_mt": _volume_mt(total_volume),
        }

    q = db.session.query(
        Municipality.id,
        Municipality.name,
        func.sum(ProductionRecord.production_volume).label("total_volume"),
        func.count(ProductionRecord.id).label("record_count"),
        func.sum(ProductionRecord.num_salt_beds).label("total_beds"),
        func.sum(ProductionRecord.num_salt_beds * ProductionRecord.area_per_salt_bed).label("total_area"),
    ).join(ProductionRecord, ProductionRecord.municipality_id == Municipality.id)
    if start:
        q = q.filter(ProductionRecord.record_date >= start)
    if end:
        q = q.filter(ProductionRecord.record_date <= end)
    rows = q.group_by(Municipality.id, Municipality.name).all()

    total_volume = sum(float(r.total_volume or 0) for r in rows)
    labels, values, _ = _monthly_aggregates(None, start, end) if start and end else ([], [], [])

    return {
        "report_title": "Provincial Production Summary",
        "by_municipality": [
            {
                "municipality": r.name,
                "production_kg": round(float(r.total_volume or 0), 2),
                "production_mt": _volume_mt(r.total_volume),
                "record_count": r.record_count,
                "total_salt_beds": int(r.total_beds or 0),
                "total_area_sqm": round(float(r.total_area or 0), 2),
            }
            for r in rows
        ],
        "monthly_trend": [
            {"month": labels[i], "production_mt": round(values[i] / 1000, 2)}
            for i in range(len(labels))
        ],
        "total_production_mt": _volume_mt(total_volume),
    }


def build_producers_data(municipality_id, barangay_id, start, end):
    """Producer-focused report across province, municipality, or barangay scope."""
    producer_cols = (
        func.sum(ProductionRecord.registered_producers).label("total_registered"),
        func.sum(ProductionRecord.male_producers).label("total_male"),
        func.sum(ProductionRecord.female_producers).label("total_female"),
        func.count(ProductionRecord.id).label("record_count"),
        func.sum(ProductionRecord.production_volume).label("total_volume"),
    )

    def apply_date(q):
        if start:
            q = q.filter(ProductionRecord.record_date >= start)
        if end:
            q = q.filter(ProductionRecord.record_date <= end)
        return q

    if barangay_id is not None:
        barangay = db.session.get(Barangay, barangay_id)
        prefix = barangay.name if barangay else f"Barangay #{barangay_id}"
        muni_name = barangay.municipality.name if barangay and barangay.municipality else None

        q = apply_date(db.session.query(Barangay.name.label("barangay"), *producer_cols)
                       .join(ProductionRecord, ProductionRecord.barangay_id == Barangay.id)
                       .filter(ProductionRecord.barangay_id == barangay_id))
        rows = q.group_by(Barangay.name).all()
        report_title = f"{prefix} Producers Report"
        table_key = "by_barangay"
        level_label = "barangay"
    elif municipality_id is not None:
        muni = db.session.get(Municipality, municipality_id)
        prefix = muni.name if muni else f"Municipality #{municipality_id}"

        q = apply_date(db.session.query(Barangay.name.label("barangay"), *producer_cols)
                       .join(ProductionRecord, ProductionRecord.barangay_id == Barangay.id)
                       .filter(ProductionRecord.municipality_id == municipality_id))
        rows = q.group_by(Barangay.name).all()
        report_title = f"{prefix} Producers Report"
        table_key = "by_barangay"
        level_label = "barangay"
        muni_name = muni.name if muni else None
    else:
        q = apply_date(db.session.query(Municipality.name.label("municipality"), *producer_cols)
                       .join(ProductionRecord, ProductionRecord.municipality_id == Municipality.id))
        rows = q.group_by(Municipality.name).all()
        report_title = "Provincial Producers Report"
        table_key = "by_municipality"
        level_label = "municipality"
        muni_name = None

    table_rows = []
    for r in rows:
        row = {
            level_label: r[0],
            "record_count": r.record_count,
            "registered_producers": int(r.total_registered or 0),
            "male_producers": int(r.total_male or 0),
            "female_producers": int(r.total_female or 0),
            "production_mt": _volume_mt(r.total_volume),
        }
        table_rows.append(row)

    total_registered = sum(row["registered_producers"] for row in table_rows)
    total_male = sum(row["male_producers"] for row in table_rows)
    total_female = sum(row["female_producers"] for row in table_rows)

    return {
        "report_title": report_title,
        "municipality_name": muni_name,
        table_key: table_rows,
        "total_registered_producers": total_registered,
        "total_male_producers": total_male,
        "total_female_producers": total_female,
    }


def build_forecast_data():
    outlook = []
    runs = (
        ForecastRun.query.options(joinedload(ForecastRun.municipality))
        .order_by(ForecastRun.created_at.desc())
        .all()
    )
    runs_data = [
        {
            "municipality": r.municipality.name if r.municipality else "Province-wide",
            "algorithm": r.algorithm,
            "readiness": r.readiness,
            "reliability": r.reliability,
            "trend_direction": r.trend_direction,
            "expected_change_pct": float(r.expected_change_pct) if r.expected_change_pct is not None else None,
            "projected_total_mt": round(float(r.projected_total or 0) / 1000, 2) if r.projected_total is not None else None,
            "mae": float(r.mae) if r.mae is not None else None,
            "rmse": float(r.rmse) if r.rmse is not None else None,
            "mape": float(r.mape) if r.mape is not None else None,
            "r2": float(r.r2) if r.r2 is not None else None,
            "period": f"{r.period_start.isoformat()} to {r.period_end.isoformat()}",
        }
        for r in runs
    ]

    munis = Municipality.query.order_by(Municipality.name).all()
    for m in munis:
        outlook.append({
            "municipality": m.name,
            "readiness": "",
            "trend_direction": "",
            "expected_change_pct": None,
        })

    return {
        "report_title": "Forecast Report",
        "runs": runs_data,
    }


def build_data_quality_data():
    records = ProductionRecord.query.all()
    total = len(records)
    WEIGHTS = {
        "production_volume": 0.20,
        "num_salt_beds": 0.10,
        "area_per_salt_bed": 0.10,
        "registered_producers": 0.15,
        "male_producers": 0.10,
        "female_producers": 0.10,
        "record_date": 0.05,
        "production_method": 0.05,
        "barangay_id": 0.15,
    }
    FIELDS = list(WEIGHTS.keys())
    by_muni = {}
    for r in records:
        bucket = by_muni.setdefault(r.municipality_id, {
            "count": 0,
            "present": {f: 0 for f in FIELDS},
            "name": r.municipality.name if r.municipality else "Unknown",
        })
        bucket["count"] += 1
        for f in FIELDS:
            v = getattr(r, f, None)
            if f == "area_per_salt_bed":
                present = v is not None
            else:
                present = v is not None and v != ""
            if present:
                bucket["present"][f] += 1

    muni_scores = []
    overall_present = {f: 0 for f in FIELDS}
    for bid, b in by_muni.items():
        score = 0.0
        completeness = {}
        for f in FIELDS:
            rate = (b["present"][f] / b["count"]) if b["count"] else 0.0
            completeness[f] = round(rate * 100, 1)
            score += rate * WEIGHTS[f]
            overall_present[f] += b["present"][f]
        muni_scores.append({
            "municipality": b["name"],
            "record_count": b["count"],
            "quality_score": round(score * 100, 1),
            "completeness": completeness,
        })
    muni_scores.sort(key=lambda x: x["quality_score"])

    overall = 0.0
    if total:
        for f in FIELDS:
            overall += (overall_present[f] / total) * WEIGHTS[f]

    return {
        "report_title": "Data Quality Report",
        "overall_quality_score": round(overall * 100, 1),
        "municipalities": muni_scores,
        "total_records": total,
    }


def build_supply_demand_data():
    provincial = DemandBenchmark.query.filter_by(geographic_scope="provincial").order_by(DemandBenchmark.year.desc()).first()
    national = DemandBenchmark.query.filter_by(geographic_scope="national").order_by(DemandBenchmark.year.desc()).first()
    local_kg = db.session.query(func.sum(ProductionRecord.production_volume)).filter(ProductionRecord.status == "approved").scalar() or 0
    local_mt = round(local_kg / 1000, 2)

    def scope(row):
        if row is None:
            return {}
        return {
            "year": row.year,
            "demand_volume_mt": float(row.demand_volume),
            "local_production_mt": float(row.local_production) if row.local_production is not None else None,
            "import_volume_mt": float(row.import_volume) if row.import_volume is not None else None,
            "source": row.source_name,
        }

    return {
        "report_title": "Supply & Demand Report",
        "pangasinan": {
            "year": provincial.year if provincial else datetime.utcnow().year,
            "demand_volume_mt": float(provincial.demand_volume) if provincial else None,
            "local_production_mt": local_mt if local_mt > 0 else (float(provincial.local_production) if provincial and provincial.local_production else None),
            "source": provincial.source_name if provincial else None,
        },
        "philippines": scope(national),
        "sector_demand": SECTOR_DEMAND,
    }


def build_gis_data():
    munis = Municipality.query.order_by(Municipality.name).all()
    record_counts = dict(
        db.session.query(ProductionRecord.municipality_id, func.count(ProductionRecord.id))
        .group_by(ProductionRecord.municipality_id)
        .all()
    )
    return {
        "report_title": "Geographic Reference Report",
        "municipalities": [
            {
                "municipality": m.name,
                "latitude": float(m.latitude) if m.latitude is not None else None,
                "longitude": float(m.longitude) if m.longitude is not None else None,
                "record_count": record_counts.get(m.id, 0),
            }
            for m in munis
        ],
    }


BUILDERS = {
    "provincial": build_provincial_data,
    "municipality": build_municipality_data,
    "forecast": build_forecast_data,
    "data_quality": build_data_quality_data,
    "supply_demand": build_supply_demand_data,
    "gis": build_gis_data,
    "production": build_production_data,
    "producers": build_producers_data,
}


def build_report_data(report_type, municipality_id=None, barangay_id=None, start=None, end=None):
    # WHAT: Piliin ang tamang data-builder base sa report_type.
    # WHY: Isang entry point para sa lahat ng report types, iwas if/elif sa API.
    fn = BUILDERS.get(report_type)
    if fn is None:
        raise ValueError(f"Unsupported report type: {report_type}")
    if report_type == "municipality":
        return fn(municipality_id, start, end)
    if report_type == "provincial":
        return fn(start, end)
    if report_type in ("production", "producers"):
        return fn(municipality_id, barangay_id, start, end)
    return fn()


def _flatten_rows(rows):
    flat = []
    if not rows:
        return flat
    for row in rows:
        if not isinstance(row, dict):
            continue
        out = {}
        for k, v in row.items():
            if isinstance(v, dict):
                out.update({f"{k}_{sk}": sv for sk, sv in v.items()})
            else:
                out[k] = v
        flat.append(out)
    return flat


def write_excel(data, filepath):
    from openpyxl import Workbook
    wb = Workbook()

    def sheet_from_rows(ws, rows):
        if not rows:
            ws.append(["No data"])
            return
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])

    ws = wb.active
    ws.title = "Summary"
    ws.append([data.get("report_title", "Report")])
    ws.append(["Generated", datetime.utcnow().isoformat() + "Z"])
    ws.append([])

    if "municipalities" in data and isinstance(data["municipalities"], list):
        sheet = wb.create_sheet("Municipalities")
        sheet_from_rows(sheet, _flatten_rows(data["municipalities"]))
    if "runs" in data and isinstance(data["runs"], list):
        sheet = wb.create_sheet("Forecast Runs")
        sheet_from_rows(sheet, data["runs"])
    if "by_municipality" in data and isinstance(data["by_municipality"], list):
        sheet = wb.create_sheet("By Municipality")
        sheet_from_rows(sheet, data["by_municipality"])
    if "by_barangay" in data and isinstance(data["by_barangay"], list):
        sheet = wb.create_sheet("By Barangay")
        sheet_from_rows(sheet, data["by_barangay"])
    if "monthly_trend" in data and isinstance(data["monthly_trend"], list):
        sheet = wb.create_sheet("Monthly Trend")
        sheet_from_rows(sheet, data["monthly_trend"])

    for extra in (
        "total_production_mt",
        "overall_quality_score",
        "total_records",
        "total_registered_producers",
        "total_male_producers",
        "total_female_producers",
    ):
        if extra in data and data[extra] not in (None, ""):
            ws.append([extra.replace("_", " ").title(), data[extra]])

    if "pangasinan" in data and isinstance(data["pangasinan"], dict):
        ws.append([])
        ws.append(["Pangasinan"])
        for k, v in data["pangasinan"].items():
            ws.append([k.replace("_", " ").title(), v])
    if "philippines" in data and isinstance(data["philippines"], dict):
        ws.append([])
        ws.append(["Philippines"])
        for k, v in data["philippines"].items():
            ws.append([k.replace("_", " ").title(), v])

    wb.save(filepath)


def write_pdf(data, filepath):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=16, spaceAfter=12)
    h_style = ParagraphStyle("HeadX", parent=styles["Heading2"], fontSize=11, spaceBefore=8, spaceAfter=4)
    small_style = ParagraphStyle(name="SmallX", parent=styles["BodyText"], fontSize=7)

    doc = SimpleDocTemplate(filepath, pagesize=landscape(A4))
    story = [Paragraph(data.get("report_title", "Report"), title_style)]
    story.append(Paragraph(f"Generated: {datetime.utcnow().isoformat()}Z", styles["BodyText"]))
    story.append(Spacer(1, 10))

    table_sections = [
        ("Municipalities", "municipalities"),
        ("By Municipality", "by_municipality"),
        ("By Barangay", "by_barangay"),
        ("Monthly Trend", "monthly_trend"),
        ("Forecast Runs", "runs"),
    ]
    for title, key in table_sections:
        rows = data.get(key)
        if not rows:
            continue
        story.append(Paragraph(title, h_style))
        flat = _flatten_rows(rows)
        if not flat:
            continue
        headers = list(flat[0].keys())
        body = [[Paragraph(str(v), small_style) for v in row.values()] for row in flat]
        tbl = Table([[Paragraph(h, small_style) for h in headers]] + body, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef3fb")]),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 8))

    if "overall_quality_score" in data:
        story.append(Paragraph(f"Overall Quality Score: {data['overall_quality_score']}%", styles["Normal"]))
    if "total_production_mt" in data:
        story.append(Paragraph(f"Total Production: {data['total_production_mt']} MT", styles["Normal"]))
    if "total_registered_producers" in data:
        story.append(Paragraph(f"Total Registered Producers: {data['total_registered_producers']}", styles["Normal"]))
        story.append(Paragraph(
            f"Male: {data['total_male_producers']} · Female: {data['total_female_producers']}",
            styles["Normal"],
        ))

    doc.build(story)


def write_docx(data, filepath):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    title = doc.add_heading(data.get("report_title", "Report"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    generated = doc.add_paragraph()
    run = generated.add_run(f"Generated: {datetime.utcnow().isoformat()}Z")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    def set_cell_shading(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def add_table(title, rows):
        if not rows:
            return
        doc.add_heading(title, level=1)
        headers = list(rows[0].keys())
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for j, h in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = str(h)
            set_cell_shading(cell, "1565C8")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.bold = True

        for i, row in enumerate(rows, start=1):
            for j, h in enumerate(headers):
                cell = table.cell(i, j)
                cell.text = "" if row.get(h) is None else str(row.get(h, ""))
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)

        doc.add_paragraph()

    table_sections = [
        ("Municipalities", "municipalities"),
        ("By Municipality", "by_municipality"),
        ("By Barangay", "by_barangay"),
        ("Monthly Trend", "monthly_trend"),
        ("Forecast Runs", "runs"),
    ]
    for title, key in table_sections:
        rows = data.get(key)
        if not rows:
            continue
        add_table(title, _flatten_rows(rows))

    if "overall_quality_score" in data:
        doc.add_paragraph(f"Overall Quality Score: {data['overall_quality_score']}%")
    if "total_production_mt" in data:
        doc.add_paragraph(f"Total Production: {data['total_production_mt']} MT")
    if "total_registered_producers" in data:
        doc.add_paragraph(f"Total Registered Producers: {data['total_registered_producers']}")
        doc.add_paragraph(
            f"Male: {data['total_male_producers']} · Female: {data['total_female_producers']}"
        )

    doc.save(filepath)


def generate_report_file(report_type, fmt, municipality_id=None, barangay_id=None, start=None, end=None, reports_dir=None):
    # WHAT: Buuin ang report data at isulat sa file (excel/pdf/docx).
    # WHY: Ang API ay nagko-call dito; ito ang nagde-decide ng extension at writer.
    data = build_report_data(report_type, municipality_id, barangay_id, start, end)
    if fmt == "excel":
        ext = "xlsx"
        writer = write_excel
    elif fmt == "docx":
        ext = "docx"
        writer = write_docx
    else:
        ext = "pdf"
        writer = write_pdf
    filename = f"{report_type}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.{ext}"
    filepath = os.path.join(reports_dir or "", filename)
    writer(data, filepath)
    return filename, filepath, data
